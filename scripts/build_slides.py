"""Render docs/slides/slides.html to PNGs, then build Sting-slides.pptx (speaker notes = narration) and a PDF.

  python scripts/build_slides.py
"""
import json
from pathlib import Path

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
SL = ROOT / "docs" / "slides"

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
    pg.goto((SL / "slides.html").as_uri())
    pg.wait_for_timeout(1500)
    ids = pg.eval_on_selector_all(".slide", "els => els.map(e => e.id)")
    notes = pg.eval_on_selector_all(".slide", "els => els.map(e => e.dataset.notes)")
    for i in ids:
        pg.locator(f"#{i}").screenshot(path=str(SL / f"{i}.png"))
    b.close()

(SL / "notes.json").write_text(json.dumps(dict(zip(ids, notes)), indent=1), encoding="utf-8")

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
for i, n in zip(ids, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(SL / f"{i}.png"), 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = n
cp = prs.core_properties
cp.author = cp.last_modified_by = "Paras Adlakha; Harshita Adlakha"
cp.title = "The Undercover Firm"
prs.save(SL / "Sting-slides.pptx")

imgs = [Image.open(SL / f"{i}.png").convert("RGB") for i in ids]
imgs[0].save(SL / "Sting-slides.pdf", save_all=True, append_images=imgs[1:], resolution=144)
print("slides:", len(ids))
