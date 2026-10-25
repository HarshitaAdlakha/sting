"""Build the submission report from report/report.md into the official Apart template, then PDF.

  python scripts/build_report.py            -> report/Sting-report.docx and report/Sting-report.pdf

report.md conventions:
  first lines: "title: ...", "authors: Name | Affiliation; Name | Affiliation", "abstract: ..." (one line)
  "## Heading" -> Heading 2, "### Heading" -> Heading 3, "- item" bullets, "1. item" numbered,
  "![Figure 1: caption](path){width=6.3}" figures, pipe tables (a line "Table N: caption" right above),
  **bold**, *italic*, [text](url) links, "\\pagebreak" page break.
"""
from __future__ import annotations

import copy
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT.parent / "Copy of Apart Research hackathon submission template.docx"
SRC = ROOT / "report" / "report.md"
OUT = ROOT / "report" / "Sting-report.docx"
FONT = "Old Standard TT"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|`[^`]+`)")


def add_runs(p, text, size=None, italic=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
            continue
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(p, label, url, size)
            continue
        elif part.startswith("*"):
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        r.font.name = FONT
        if size:
            r.font.size = Pt(size)
        if italic:
            r.italic = True


def add_hyperlink(p, text, url, size=None):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1155CC"), ("w:u", "single")):
        e = OxmlElement(tag); e.set(qn("w:val"), val); rpr.append(e)
    f = OxmlElement("w:rFonts"); f.set(qn("w:ascii"), FONT); f.set(qn("w:hAnsi"), FONT); rpr.append(f)
    if size:
        s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(size * 2))); rpr.append(s)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)


def set_cell_text(tc_el, lines):
    """Replace the text of a <w:tc> with given lines, keeping the first paragraph's formatting."""
    ps = tc_el.findall(qn("w:p"))
    keep = [p for p in ps if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()]
    proto = keep[0] if keep else ps[0]
    for p in ps:
        tc_el.remove(p)
    for i, line in enumerate(lines):
        p = copy.deepcopy(proto)
        runs = p.findall(qn("w:r"))
        for r in runs[1:]:
            p.remove(r)
        t = runs[0].find(qn("w:t"))
        t.text = line
        if i == 1:  # affiliation line: not bold
            rpr = runs[0].find(qn("w:rPr"))
            if rpr is not None and rpr.find(qn("w:b")) is not None:
                rpr.remove(rpr.find(qn("w:b")))
        tc_el.append(p)


def fill_header(d, title, authors, abstract):
    t = d.tables[0]._tbl
    # title
    for tnode in t.iter(qn("w:t")):
        if tnode.text and "PROJECT TITLE" in tnode.text:
            tnode.text = title
    # authors: nested table with rows of 3 cells; keep one row with len(authors) cells
    auth_tbl = [tb for tb in t.iter(qn("w:tbl")) if "Author name 1" in "".join(x.text or "" for x in tb.iter(qn("w:t")))]
    auth_tbl = min(auth_tbl, key=lambda tb: len(list(tb.iter())))
    rows = auth_tbl.findall(qn("w:tr"))
    for r in rows[1:]:
        auth_tbl.remove(r)
    cells = rows[0].findall(qn("w:tc"))
    for c in cells[len(authors):]:
        rows[0].remove(c)
    grid = auth_tbl.find(qn("w:tblGrid"))
    cols = grid.findall(qn("w:gridCol"))
    for c in cols[len(authors):]:
        grid.remove(c)
    width = 3000
    for c in grid.findall(qn("w:gridCol")):
        c.set(qn("w:w"), str(width))
    for (name, aff), c in zip(authors, rows[0].findall(qn("w:tc"))):
        tcpr = c.find(qn("w:tcPr"))
        if tcpr is not None and tcpr.find(qn("w:tcW")) is not None:
            tcpr.find(qn("w:tcW")).set(qn("w:w"), str(width)); tcpr.find(qn("w:tcW")).set(qn("w:type"), "dxa")
        ts = list(c.iter(qn("w:t")))   # template cell: "Author name N" <br/> "Affiliation" in one run
        ts[0].text, ts[1].text = name, aff
    tblw = auth_tbl.find(qn("w:tblPr")).find(qn("w:tblW"))
    if tblw is not None:
        tblw.set(qn("w:w"), str(width * len(authors))); tblw.set(qn("w:type"), "dxa")
    # abstract
    abs_tbl = [tb for tb in t.iter(qn("w:tbl")) if "Summarize your project" in "".join(x.text or "" for x in tb.iter(qn("w:t")))][-1]
    tc = abs_tbl.find(qn("w:tr")).find(qn("w:tc"))
    ps = tc.findall(qn("w:p"))
    target = [p for p in ps if "Summarize" in "".join(x.text or "" for x in p.iter(qn("w:t")))][0]
    runs = target.findall(qn("w:r"))
    for r in runs[1:]:
        target.remove(r)
    runs[0].find(qn("w:t")).text = abstract
    rpr = runs[0].find(qn("w:rPr"))
    if rpr is not None:
        for tag in ("w:i", "w:iCs", "w:u"):
            e = rpr.find(qn(tag))
            if e is not None:
                rpr.remove(e)


def clear_body_after_header(d):
    body = d.element.body
    started = False
    for el in list(body):
        if el.tag == qn("w:tbl") and not started:
            if "How to use this template" in "".join(x.text or "" for x in el.iter(qn("w:t"))):
                body.remove(el)
                started = True
            continue
        if started and el.tag != qn("w:sectPr"):
            body.remove(el)


def para(d, text, style="normal", align=None, size=None, space_after=6, italic=None):
    p = d.add_paragraph(style=style)
    add_runs(p, text, size=size, italic=italic)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align is None and style == "normal":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def list_item(d, text, numbered, n):
    p = d.add_paragraph(style="normal")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(3)
    add_runs(p, (f"{n}. " if numbered else "• ") + text)


def table(d, rows, caption):
    if caption:
        para(d, caption, align="center", size=9.5, space_after=3)
    hdr, body = rows[0], rows[1:]
    t = d.add_table(rows=len(rows), cols=len(hdr))
    t.style = d.styles["Table Grid"] if "Table Grid" in [s.name for s in d.styles] else None
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.paragraphs[0].text = ""
            add_runs(c.paragraphs[0], val.strip(), size=8.5)
            if i == 0:
                for r in c.paragraphs[0].runs:
                    r.bold = True
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    _borders(t)
    para(d, "", space_after=4)


def _borders(t):
    tblpr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
        b.append(e)
    tblpr.append(b)


def figure(d, path, caption, width):
    p = d.add_paragraph(style="normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / path), width=Inches(width))
    p.paragraph_format.space_after = Pt(2)
    para(d, caption, align="center", size=9.5, space_after=10)


def build():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    meta = {}
    while lines and re.match(r"^(title|authors|abstract):", lines[0]):
        k, v = lines.pop(0).split(":", 1)
        meta[k] = v.strip()
    authors = [tuple(x.strip() for x in a.split("|")) for a in meta["authors"].split(";")]
    d = docx.Document(TEMPLATE)
    fill_header(d, meta["title"], authors, meta["abstract"])
    clear_body_after_header(d)

    i, n_num, pending_caption = 0, 0, None
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            n_num = 0; i += 1; continue
        if ln.strip() == "\\pagebreak":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE); i += 1; continue
        if ln.startswith("### "):
            d.add_paragraph(ln[4:], style="Heading 3"); i += 1; continue
        if ln.startswith("## "):
            d.add_paragraph(ln[3:], style="Heading 2"); i += 1; continue
        m = re.match(r"!\[(.*)\]\((.*?)\)(\{width=([\d.]+)\})?", ln)
        if m:
            figure(d, m.group(2), m.group(1), float(m.group(4) or 6.3)); i += 1; continue
        if re.match(r"^Table \d+:", ln) and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            pending_caption = ln; i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*:?-+:?\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(d, rows, pending_caption); pending_caption = None; continue
        if ln.startswith("- "):
            list_item(d, ln[2:], False, 0); i += 1; continue
        m = re.match(r"^(\d+)\. (.*)", ln)
        if m:
            list_item(d, m.group(2), True, int(m.group(1))); i += 1; continue
        # paragraph: join following non-empty plain lines
        buf = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|- |\d+\. |\||!\[|\\pagebreak|Table \d+:)", lines[i]):
            buf.append(lines[i].strip()); i += 1
        size = 9 if buf[0].startswith("[") and "References" in "" else None
        para(d, " ".join(buf), size=size)
    cp = d.core_properties
    cp.author = cp.last_modified_by = "; ".join(a[0] for a in authors)
    cp.title = meta["title"]
    d.save(OUT)
    print("wrote", OUT)


def to_pdf():
    pdf = OUT.with_suffix(".pdf")
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{OUT}', $false, $true); $d.SaveAs([ref]'{pdf}', [ref]17); "
          f"$n = $d.ComputeStatistics(2); $d.Close($false); $w.Quit(); \"pages=$n\"")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    print(r.stdout.strip(), r.stderr.strip()[:500])


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        to_pdf()
